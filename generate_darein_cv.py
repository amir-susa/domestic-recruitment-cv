import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_darein_cv():
    doc = docx.Document()
    
    # Page Setup (A4 Portrait, 0.5 in margins)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    NAVY_HEX = "000080"
    NAVY_RGB = RGBColor(0, 0, 128)
    LIGHT_BLUE_HEX = "8DB3E2"
    WHITE_RGB = RGBColor(255, 255, 255)

    def set_cell_background(cell, fill_hex):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    # Header
    header = doc.sections[0].header
    header_table = header.add_table(rows=3, cols=3, width=Inches(7.27))
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell_l0 = header_table.cell(0, 0)
    p_l0 = cell_l0.paragraphs[0]
    r_logo = p_l0.add_run("D ")
    r_logo.font.name = 'Garamond'
    r_logo.font.size = Pt(16)
    r_logo.font.bold = True
    r_logo.font.color.rgb = NAVY_RGB
    r_name = p_l0.add_run("DAREIN EST")
    r_name.font.name = 'Garamond'
    r_name.font.size = Pt(14)
    r_name.font.bold = True
    r_name.font.color.rgb = NAVY_RGB

    cell_r0 = header_table.cell(0, 2)
    p_r0 = cell_r0.paragraphs[0]
    p_r0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ar_name = p_r0.add_run("دارين التأسيس")
    r_ar_name.font.name = 'Arial'
    r_ar_name.font.size = Pt(14)
    r_ar_name.font.bold = True
    r_ar_name.font.color.rgb = NAVY_RGB
    r_ar_name.font.rtl = True

    # Address Row
    cell_addr = header_table.cell(2, 0)
    cell_addr.merge(header_table.cell(2, 1)).merge(header_table.cell(2, 2))
    set_cell_background(cell_addr, NAVY_HEX)
    p_addr = cell_addr.paragraphs[0]
    p_addr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_ad1 = p_addr.add_run("الفروانية بلوك ٤ حبيب منور شارع - عريد معرض - مكتب الميزانين رقم ١ - الكويت\n")
    r_ad1.font.name = 'Arial'
    r_ad1.font.size = Pt(8)
    r_ad1.font.bold = True
    r_ad1.font.color.rgb = WHITE_RGB
    r_ad1.font.rtl = True

    r_ad2 = p_addr.add_run("Farwaniya Block 4, Habib Menawer Street, Arbeed gallery, Mezanin office No. 1, Kuwait")
    r_ad2.font.name = 'Garamond'
    r_ad2.font.size = Pt(8)
    r_ad2.font.bold = True
    r_ad2.font.color.rgb = WHITE_RGB

    # Footer
    footer = doc.sections[0].footer
    footer_table = footer.add_table(rows=1, cols=2, width=Inches(7.27))
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    f_cell_l = footer_table.cell(0, 0)
    set_cell_background(f_cell_l, NAVY_HEX)
    p_fl = f_cell_l.paragraphs[0]
    r_fl = p_fl.add_run("DAREIN EST – Application for Employment")
    r_fl.font.name = 'Garamond'
    r_fl.font.size = Pt(8.5)
    r_fl.font.bold = True
    r_fl.font.color.rgb = WHITE_RGB

    f_cell_r = footer_table.cell(0, 1)
    set_cell_background(f_cell_r, NAVY_HEX)
    p_fr = f_cell_r.paragraphs[0]
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_fr = p_fr.add_run("Date / التاريخ: ______________   |   Ref. No / رقم المرجع: ______________")
    r_fr.font.name = 'Garamond'
    r_fr.font.size = Pt(8)
    r_fr.font.bold = True
    r_fr.font.color.rgb = WHITE_RGB

    # Body Title
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_t.add_run("Application for Employment")
    r_t.font.name = 'Garamond'
    r_t.font.size = Pt(14)
    r_t.font.bold = True
    r_t.font.color.rgb = NAVY_RGB

    doc.save("Darein_Est_Application_For_Employment.docx")
    print("Document successfully created: Darein_Est_Application_For_Employment.docx")

if __name__ == "__main__":
    create_darein_cv()